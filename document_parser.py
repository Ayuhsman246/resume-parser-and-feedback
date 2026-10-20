"""Document text extraction module.

Handles extracting text from PDF, DOCX, and TXT file uploads.
Each format has a dedicated extractor function.
"""

import io
import os

import docx  # provided by the python-docx package
import PyPDF2

SUPPORTED_TYPES = ["pdf", "docx", "txt"]


def _extract_pdf(file_bytes):
    """Extract text from PDF bytes using PyPDF2."""
    reader = PyPDF2.PdfReader(io.BytesIO(file_bytes))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(pages)


def _extract_docx(file_bytes):
    """Extract text from DOCX bytes using python-docx (paragraphs + tables)."""
    document = docx.Document(io.BytesIO(file_bytes))
    lines = [para.text for para in document.paragraphs]

    # Resumes often use tables for layout, so include their cell text too.
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    lines.append(cell.text)
    return "\n".join(lines)


def _extract_txt(file_bytes):
    """Decode plain-text bytes, falling back to latin-1 if UTF-8 fails."""
    try:
        return file_bytes.decode("utf-8")
    except UnicodeDecodeError:
        return file_bytes.decode("latin-1")


# Map file extensions to their extractor functions.
_EXTRACTORS = {
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
    ".txt": _extract_txt,
}


def extract_text(uploaded_file):
    """Detect the file type from its extension and extract its text.

    Args:
        uploaded_file: Streamlit UploadedFile object (PDF, DOCX or TXT).

    Returns:
        str: The extracted text.

    Raises:
        ValueError: If the type is unsupported or no text could be extracted.
    """
    extension = os.path.splitext(uploaded_file.name)[1].lower()
    file_bytes = uploaded_file.getvalue()

    if extension not in _EXTRACTORS:
        raise ValueError(f"Unsupported file type: '{extension}'.")

    try:
        text = _EXTRACTORS[extension](file_bytes)
    except Exception as exc:  # corrupt or encrypted files, etc.
        raise ValueError(f"Could not read the {extension} file: {exc}") from exc

    if not text.strip():
        raise ValueError(
            "No text could be extracted. The file may be empty or a scanned image."
        )
    return text

